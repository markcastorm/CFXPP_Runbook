"""
CFXPP Comparison Tool V2
========================
Dynamic path discovery - no manual path entry needed.

What it auto-discovers:
  - Pipeline output:    Latest run in ../output/  (CFXPP_DATA_*.xlsx)
  - Source mapping:     source_column_mapping.csv from same output run
  - Archive directory:  Latest folder in ../archive/
  - Reference file:     Any .xlsx dropped into ./reference_input/

Usage:
  1. Drop your reference .xlsx into the  reference_input/  folder
  2. Run: python compare_v2.py
  3. Find results in the new run_YYYYMMDD_HHMMSS/ folder

Settings (tolerances, verbosity) are in config_v2.py.
"""

import openpyxl
from openpyxl.styles import PatternFill, Font
import csv
import os
import sys
import glob
import shutil
from datetime import datetime
from collections import defaultdict

import config_v2 as config


# ---------------------------------------------------------------------------
# Path Discovery
# ---------------------------------------------------------------------------

def _sorted_subdirs(parent):
    """Return subdirectories of *parent* sorted ascending (oldest first)."""
    if not os.path.isdir(parent):
        return []
    entries = [
        e for e in os.listdir(parent)
        if os.path.isdir(os.path.join(parent, e))
    ]
    return sorted(entries)


def find_latest_output_run(project_root):
    """
    Locate the most-recent timestamped run folder under  <project_root>/output/
    and return (data_xlsx_path, mapping_csv_path).

    Only considers folders that match the YYYYMMDD_HHMMSS timestamp pattern
    so that alias folders like 'latest' are ignored.
    """
    import re
    output_root = os.path.join(project_root, 'output')
    subdirs = _sorted_subdirs(output_root)

    # Filter to timestamp-style folders only
    ts_pattern = re.compile(r'^\d{8}_\d{6}$')
    ts_subdirs = [s for s in subdirs if ts_pattern.match(s)]

    if not ts_subdirs:
        raise FileNotFoundError(f'No timestamped run folders (YYYYMMDD_HHMMSS) found in: {output_root}')

    # Walk from newest to oldest until we find a folder with a DATA xlsx
    for folder_name in reversed(ts_subdirs):
        folder = os.path.join(output_root, folder_name)
        data_files = glob.glob(os.path.join(folder, 'CFXPP_DATA_*.xlsx'))
        if data_files:
            data_xlsx = data_files[0]
            mapping_csv = os.path.join(folder, 'source_column_mapping.csv')
            return data_xlsx, mapping_csv

    raise FileNotFoundError(
        f'No CFXPP_DATA_*.xlsx found in any output run folder under: {output_root}'
    )


def find_latest_archive_dir(project_root):
    """
    Return the path to the most-recent (non-zip) folder under
    <project_root>/archive/.  Returns '' if none found.
    """
    archive_root = os.path.join(project_root, 'archive')
    # Only real directories (skip .zip files)
    subdirs = _sorted_subdirs(archive_root)
    if not subdirs:
        return ''
    return os.path.join(archive_root, subdirs[-1])


def find_reference_file(compare_v2_dir):
    """
    Return the path to the single .xlsx in reference_input/.
    Raises clear errors if 0 or >1 files are found.
    """
    ref_dir = os.path.join(compare_v2_dir, 'reference_input')
    os.makedirs(ref_dir, exist_ok=True)

    xlsx_files = glob.glob(os.path.join(ref_dir, '*.xlsx'))
    # Filter out temp Excel lock files (~$...)
    xlsx_files = [f for f in xlsx_files if not os.path.basename(f).startswith('~$')]

    if not xlsx_files:
        raise FileNotFoundError(
            f'No reference .xlsx found in:\n  {ref_dir}\n\n'
            'Drop your reference file there and re-run.'
        )

    if len(xlsx_files) > 1:
        names = '\n  '.join(os.path.basename(f) for f in xlsx_files)
        raise ValueError(
            f'Multiple .xlsx files found in reference_input/ - keep only one:\n  {names}'
        )

    return xlsx_files[0]


# ---------------------------------------------------------------------------
# Comparison Engine  (same logic as V1)
# ---------------------------------------------------------------------------

class ComparisonReport:
    """Handles comparison between pipeline output and reference data."""

    def __init__(self, output_data, reference_data, mapping_csv, archive_dir):
        self.output_data_path = output_data
        self.reference_data_path = reference_data
        self.mapping_csv_path = mapping_csv
        self.archive_dir = archive_dir

        self.output_codes = []
        self.output_cells = {}
        self.reference_codes = []
        self.reference_cells = {}
        self.code_to_entries = {}
        self.descriptions = {}

        self.matches = []
        self.mismatches = []
        self.missing = []
        self.extra = []

        self.mismatch_coords = []
        self.missing_coords = []
        self.extra_coords = []

    # ------------------------------------------------------------------
    # Loaders
    # ------------------------------------------------------------------

    def load_excel_data(self, path, label):
        """Load column codes and cell values from DATA xlsx."""
        if config.VERBOSE:
            print(f'Loading {label}: {os.path.basename(path)}')

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb['DATA']

        codes = []
        for c in range(2, 542):
            code = ws.cell(row=1, column=c).value
            codes.append(code)

        cells = {}
        for r in range(2, ws.max_row + 1):
            date_val = ws.cell(row=r, column=1).value
            if date_val is None:
                continue
            d = date_val.strftime('%Y-%m-%d') if hasattr(date_val, 'strftime') else str(date_val)
            for c in range(2, 542):
                v = ws.cell(row=r, column=c).value
                if v is not None:
                    cells[(d, c - 2)] = v

        wb.close()

        if config.VERBOSE:
            print(f'  Loaded: {len(codes)} codes, {len(cells)} cells')

        return codes, cells

    def load_mapping_csv(self, path):
        """Load source-column mapping CSV."""
        if config.VERBOSE:
            print(f'Loading source mapping: {os.path.basename(path)}')

        code_to_entries = defaultdict(list)

        if not os.path.exists(path):
            print(f'  WARNING: Mapping CSV not found: {path}')
            return code_to_entries

        with open(path, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                code_to_entries[row['Output_Column_Code']].append(row)

        if config.VERBOSE:
            print(f'  Loaded: {len(code_to_entries)} column mappings')

        return code_to_entries

    def load_meta_descriptions(self, output_data_path):
        """Load column descriptions from META file."""
        meta_path = output_data_path.replace('DATA', 'META')
        descs = {}

        if not os.path.exists(meta_path):
            if config.VERBOSE:
                print(f'  No META file found: {os.path.basename(meta_path)}')
            return descs

        if config.VERBOSE:
            print(f'Loading META descriptions: {os.path.basename(meta_path)}')

        wb = openpyxl.load_workbook(meta_path, data_only=True)
        ws = wb.active

        for r in range(2, ws.max_row + 1):
            code = ws.cell(row=r, column=1).value
            desc = ws.cell(row=r, column=2).value
            if code:
                descs[code] = desc or ''

        wb.close()

        if config.VERBOSE:
            print(f'  Loaded: {len(descs)} descriptions')

        return descs

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    def get_section(self, code):
        if not code:
            return 'UNKNOWN'
        if '.CURRENCYPOSITIONING.' in code and '.G10.' in code:
            return 'G10 Currency Positioning'
        elif '.CURRENCYPOSITIONING.' in code and '.EM.' in code:
            return 'EM Currency Positioning'
        elif '.FXPAIRPOSITIONING.' in code:
            return 'FX Pair Positioning'
        return 'OTHER'

    def parse_column_metadata(self, code):
        if not code:
            return '', '', ''
        parts = code.split('.')
        if '.FXPAIRPOSITIONING.' in code:
            pair = parts[-2] if len(parts) >= 2 else ''
            if 'VOLUME_NORMALIZED' in code:
                metric = 'Volume (normalized)'
            elif 'CLOSING_PRICE' in code:
                metric = 'Closing Price'
            else:
                metric = ''
            if 'BANKS_BROKERVOLUME_NORMALIZED' in code:
                client = 'BANKS_BROKER'
            else:
                client = parts[3] if len(parts) > 3 else ''
            return pair, client, metric
        elif '.CURRENCYPOSITIONING.' in code:
            ccy = parts[-2] if len(parts) >= 2 else ''
            group = parts[-3] if len(parts) >= 3 else ''
            client = parts[-4] if len(parts) >= 4 else ''
            return f'{group}/{ccy}', client, 'Net Cumulative Positioning'
        return '', '', ''

    def compare_values(self, val1, val2):
        if val1 is None and val2 is None:
            return True, 0
        if val1 is None or val2 is None:
            return False, None
        try:
            f1, f2 = float(val1), float(val2)
            diff = f1 - f2
            if abs(diff) < config.FLOAT_TOLERANCE:
                return True, 0
            return False, diff
        except (ValueError, TypeError):
            return str(val1) == str(val2), None

    def build_row_dict(self, status, date, col_idx, code, out_val, ref_val, diff):
        section = self.get_section(code)
        pair, client, metric = self.parse_column_metadata(code)
        desc = self.descriptions.get(code, '')

        entries = self.code_to_entries.get(code, [])
        source_files, archive_links, raw_clients = [], [], []

        for e in entries:
            fn = e.get('Source_File', '')
            if fn and fn not in source_files:
                source_files.append(fn)
                if self.archive_dir:
                    archive_links.append(os.path.join(self.archive_dir, fn))
            raw = e.get('Raw_Client_Types', '')
            if raw and raw not in raw_clients:
                raw_clients.append(raw)

        return {
            'Status': status,
            'Date': date,
            'Output_Column': col_idx + 2,
            'Column_Code': code or '',
            'Section': section,
            'Currency_Pair': pair,
            'Client_Type': client,
            'Metric': metric,
            'Column_Description': desc,
            'Pipeline_Value': out_val if out_val is not None else '',
            'Reference_Value': ref_val if ref_val is not None else '',
            'Difference': round(diff, 4) if diff is not None else '',
            'Source_File(s)': '; '.join(source_files) if source_files else 'NO SOURCE FILE',
            'Archive_Link(s)': '; '.join(archive_links) if archive_links else '',
            'Raw_Client_Types': '; '.join(raw_clients) if raw_clients else '',
        }

    # ------------------------------------------------------------------
    # Core comparison
    # ------------------------------------------------------------------

    def perform_comparison(self):
        if config.VERBOSE:
            print('\nPerforming comparison...')

        all_keys = set(self.output_cells.keys()) | set(self.reference_cells.keys())

        if config.VERBOSE:
            print(f'  Total cell locations: {len(all_keys)}')

        for key in sorted(all_keys):
            date, col_idx = key
            out_val = self.output_cells.get(key)
            ref_val = self.reference_cells.get(key)
            code = self.output_codes[col_idx] if col_idx < len(self.output_codes) else ''

            if out_val is None and ref_val is None:
                continue
            elif out_val is not None and ref_val is not None:
                match, diff = self.compare_values(out_val, ref_val)
                if match:
                    self.matches.append(key)
                else:
                    self.mismatches.append(self.build_row_dict('MISMATCH', date, col_idx, code, out_val, ref_val, diff))
                    self.mismatch_coords.append(key)
            elif out_val is None:
                self.missing.append(self.build_row_dict('MISSING', date, col_idx, code, None, ref_val, None))
                self.missing_coords.append(key)
            else:
                self.extra.append(self.build_row_dict('EXTRA', date, col_idx, code, out_val, None, None))
                self.extra_coords.append(key)

        if config.VERBOSE:
            print(f'  Matches:    {len(self.matches)}')
            print(f'  Mismatches: {len(self.mismatches)}')
            print(f'  Missing:    {len(self.missing)}')
            print(f'  Extra:      {len(self.extra)}')

    # ------------------------------------------------------------------
    # Output writers
    # ------------------------------------------------------------------

    def write_csv_report(self, output_path):
        if config.VERBOSE:
            print(f'\nWriting CSV report: {os.path.basename(output_path)}')

        header = [
            'Status', 'Date', 'Output_Column', 'Column_Code',
            'Section', 'Currency_Pair', 'Client_Type', 'Metric',
            'Column_Description',
            'Pipeline_Value', 'Reference_Value', 'Difference',
            'Source_File(s)', 'Archive_Link(s)', 'Raw_Client_Types',
        ]
        all_rows = self.mismatches + self.missing + self.extra

        with open(output_path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=header)
            writer.writeheader()
            for row in all_rows:
                writer.writerow(row)

        if config.VERBOSE:
            print(f'  Written {len(all_rows)} rows')

    def generate_summary_stats(self):
        total = len(self.matches) + len(self.mismatches) + len(self.missing) + len(self.extra)
        return {
            'timestamp': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'output_file': self.output_data_path,
            'reference_file': self.reference_data_path,
            'total_cells_compared': total,
            'matches': len(self.matches),
            'mismatches': len(self.mismatches),
            'missing': len(self.missing),
            'extra': len(self.extra),
            'match_rate': (len(self.matches) / total * 100) if total > 0 else 0,
        }

    def analyze_by_category(self, rows, key_func):
        groups = defaultdict(list)
        for row in rows:
            groups[key_func(row)].append(row)
        return groups

    def write_summary_report(self, output_path):
        if config.VERBOSE:
            print(f'\nWriting summary report: {os.path.basename(output_path)}')

        stats = self.generate_summary_stats()

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('=' * 70 + '\n')
            f.write('CFXPP COMPARISON REPORT V2\n')
            f.write('=' * 70 + '\n\n')
            f.write(f"Generated: {stats['timestamp']}\n\n")

            f.write('FILES COMPARED:\n')
            f.write(f"  Pipeline Output: {os.path.basename(stats['output_file'])}\n")
            f.write(f"  Reference:       {os.path.basename(stats['reference_file'])}\n")
            f.write(f"  Archive Dir:     {os.path.basename(self.archive_dir) if self.archive_dir else 'N/A'}\n\n")

            f.write('OVERALL STATISTICS:\n')
            f.write(f"  Total cells compared: {stats['total_cells_compared']:,}\n")
            f.write(f"  Matches:              {stats['matches']:,} ({stats['match_rate']:.2f}%)\n")
            f.write(f"  Mismatches:           {stats['mismatches']:,}\n")
            f.write(f"  Missing (in ref):     {stats['missing']:,}\n")
            f.write(f"  Extra (in pipeline):  {stats['extra']:,}\n\n")

            if self.mismatches:
                f.write('MISMATCHES BY SECTION:\n')
                by_section = self.analyze_by_category(self.mismatches, lambda r: r['Section'])
                for section, items in sorted(by_section.items()):
                    f.write(f"  {section}: {len(items)}\n")
                f.write('\n')

                f.write('TOP MISMATCHES (by absolute difference):\n')
                sorted_mm = sorted(
                    [r for r in self.mismatches if r['Difference'] != ''],
                    key=lambda r: abs(float(r['Difference'])) if r['Difference'] else 0,
                    reverse=True
                )[:config.MAX_CONSOLE_EXAMPLES]
                for r in sorted_mm:
                    f.write(f"  {r['Date']} | {r['Currency_Pair']} {r['Client_Type']} {r['Metric']}\n")
                    f.write(f"    Pipeline: {r['Pipeline_Value']} | Reference: {r['Reference_Value']} | Diff: {r['Difference']}\n")
                    f.write(f"    Source: {r['Source_File(s)']}\n")
                f.write('\n')

            if self.missing:
                f.write('MISSING CELLS BY CURRENCY PAIR:\n')
                by_pair = self.analyze_by_category(self.missing, lambda r: r['Currency_Pair'])
                for pair, items in sorted(by_pair.items(), key=lambda x: -len(x[1])):
                    has_source = any(r['Source_File(s)'] != 'NO SOURCE FILE' for r in items)
                    src_note = '' if has_source else ' [NO SOURCE FILES]'
                    f.write(f"  {pair}: {len(items)} cells{src_note}\n")
                f.write('\n')

            if self.extra:
                f.write('EXTRA CELLS BY CURRENCY PAIR:\n')
                by_pair = self.analyze_by_category(self.extra, lambda r: r['Currency_Pair'])
                for pair, items in sorted(by_pair.items(), key=lambda x: -len(x[1])):
                    f.write(f"  {pair}: {len(items)} cells\n")
                f.write('\n')

            f.write('=' * 70 + '\n')
            f.write(f"Detailed CSV report: {config.REPORT_FILENAME}\n")
            f.write('=' * 70 + '\n')

    def create_annotated_excel(self, source_path, output_path, coords_dict, label):
        if config.VERBOSE:
            print(f'\nCreating annotated {label} file: {os.path.basename(output_path)}')

        shutil.copy2(source_path, output_path)
        wb = openpyxl.load_workbook(output_path)
        ws = wb['DATA']

        red_fill    = PatternFill(start_color='FFCCCC', end_color='FFCCCC', fill_type='solid')
        yellow_fill = PatternFill(start_color='FFFFCC', end_color='FFFFCC', fill_type='solid')
        green_fill  = PatternFill(start_color='CCFFCC', end_color='CCFFCC', fill_type='solid')
        bold_font   = Font(bold=True)

        date_to_row = {}
        for r in range(2, ws.max_row + 1):
            date_val = ws.cell(row=r, column=1).value
            if date_val:
                d = date_val.strftime('%Y-%m-%d') if hasattr(date_val, 'strftime') else str(date_val)
                date_to_row[d] = r

        highlight_counts = {'MISMATCH': 0, 'MISSING': 0, 'EXTRA': 0}
        fill_map = {'MISMATCH': red_fill, 'MISSING': yellow_fill, 'EXTRA': green_fill}

        for status, coords in coords_dict.items():
            fill = fill_map[status]
            for date, col_idx in coords:
                row_num = date_to_row.get(date)
                if row_num:
                    ws.cell(row=row_num, column=col_idx + 2).fill = fill
                    highlight_counts[status] += 1

        ws.insert_rows(1)
        ws.cell(row=1, column=1, value='LEGEND:').font = bold_font
        ws.cell(row=1, column=2, value='RED = Mismatch').fill = red_fill
        ws.cell(row=1, column=3, value='YELLOW = Missing').fill = yellow_fill
        ws.cell(row=1, column=4, value='GREEN = Extra').fill = green_fill
        ws.cell(row=1, column=5, value=f'Annotated: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')

        wb.save(output_path)
        wb.close()

        if config.VERBOSE:
            print(f'  Highlighted: {highlight_counts["MISMATCH"]} mismatches (red), '
                  f'{highlight_counts["MISSING"]} missing (yellow), '
                  f'{highlight_counts["EXTRA"]} extra (green)')

    # ------------------------------------------------------------------
    # Run
    # ------------------------------------------------------------------

    def run(self):
        print('=' * 70)
        print('CFXPP COMPARISON TOOL  V2  (dynamic paths)')
        print('=' * 70)
        print()
        print(f'Pipeline Output:  {self.output_data_path}')
        print(f'Reference File:   {self.reference_data_path}')
        print(f'Source Mapping:   {self.mapping_csv_path}')
        print(f'Archive Dir:      {self.archive_dir or "(not found)"}')
        print()

        # Validate
        for label, path in [('Output', self.output_data_path), ('Reference', self.reference_data_path)]:
            if not os.path.exists(path):
                print(f'ERROR: {label} file not found: {path}')
                return 1

        # Load
        self.output_codes, self.output_cells = self.load_excel_data(self.output_data_path, 'Pipeline Output')
        self.reference_codes, self.reference_cells = self.load_excel_data(self.reference_data_path, 'Reference')

        if self.output_codes != self.reference_codes:
            print('\nWARNING: Column codes differ between output and reference!')

        self.code_to_entries = self.load_mapping_csv(self.mapping_csv_path)
        self.descriptions = self.load_meta_descriptions(self.output_data_path)

        # Compare
        self.perform_comparison()

        # Build output directory
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        compare_v2_dir = os.path.dirname(os.path.abspath(__file__))
        run_dir = os.path.join(compare_v2_dir, f'run_{timestamp}')
        os.makedirs(run_dir, exist_ok=True)

        if config.VERBOSE:
            print(f'\nOutput directory: {run_dir}')

        report_csv  = os.path.join(run_dir, config.REPORT_FILENAME)
        summary_txt = os.path.join(run_dir, config.SUMMARY_FILENAME)

        # Copy originals
        if config.VERBOSE:
            print('\nCopying original files...')
        output_copy = os.path.join(run_dir, f'ORIGINAL_{os.path.basename(self.output_data_path)}')
        ref_copy    = os.path.join(run_dir, f'ORIGINAL_{os.path.basename(self.reference_data_path)}')
        shutil.copy2(self.output_data_path, output_copy)
        shutil.copy2(self.reference_data_path, ref_copy)

        # Annotated Excel
        annotated_output = os.path.join(run_dir, f'ANNOTATED_Pipeline_{os.path.basename(self.output_data_path)}')
        annotated_ref    = os.path.join(run_dir, f'ANNOTATED_Reference_{os.path.basename(self.reference_data_path)}')

        self.create_annotated_excel(
            self.output_data_path, annotated_output,
            {'MISMATCH': self.mismatch_coords, 'EXTRA': self.extra_coords},
            'Pipeline'
        )
        self.create_annotated_excel(
            self.reference_data_path, annotated_ref,
            {'MISMATCH': self.mismatch_coords, 'MISSING': self.missing_coords},
            'Reference'
        )

        # Reports
        self.write_csv_report(report_csv)
        self.write_summary_report(summary_txt)

        # Console summary
        stats = self.generate_summary_stats()
        print('\n' + '=' * 70)
        print('COMPARISON COMPLETE')
        print('=' * 70)
        print(f"\nMatches:    {stats['matches']:,} ({stats['match_rate']:.2f}%)")
        print(f"Mismatches: {stats['mismatches']:,}")
        print(f"Missing:    {stats['missing']:,}")
        print(f"Extra:      {stats['extra']:,}")
        print(f'\nOutput directory: {run_dir}')
        print('\nFiles created:')
        for name in [
            os.path.basename(report_csv),
            os.path.basename(summary_txt),
            os.path.basename(output_copy),
            os.path.basename(ref_copy),
            os.path.basename(annotated_output),
            os.path.basename(annotated_ref),
        ]:
            print(f'  {name}')
        print('=' * 70)

        return 0


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main():
    try:
        compare_v2_dir = os.path.dirname(os.path.abspath(__file__))
        project_root   = os.path.dirname(compare_v2_dir)

        print('Discovering paths...')

        output_data, mapping_csv = find_latest_output_run(project_root)
        archive_dir              = find_latest_archive_dir(project_root)
        reference_data           = find_reference_file(compare_v2_dir)

        report = ComparisonReport(output_data, reference_data, mapping_csv, archive_dir)
        return report.run()

    except (FileNotFoundError, ValueError) as e:
        print(f'\nERROR: {e}')
        return 1

    except KeyboardInterrupt:
        print('\n\nInterrupted by user')
        return 130

    except Exception as e:
        print(f'\nERROR: {e}')
        import traceback
        traceback.print_exc()
        return 1


if __name__ == '__main__':
    sys.exit(main())
